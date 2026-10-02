"""OCR de planos vectoriales (texto dibujado como lineas) -> PDF buscable + indice Excel.

Uso: python ocr_plano.py "entrada.pdf" [carpeta_salida]
"""
import sys, os, re, io
import numpy as np, cv2, pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.colors import Color
from rapidocr import RapidOCR

SCALE = 3.0          # px por punto PDF
TILE, OVER = 1800, 300
MIN_SCORE = 0.75

# Rejilla del cajetin (fracciones del ancho/alto de la hoja)
COL_EDGES = [x / 1191 for x in (214, 346, 478, 610, 741, 873, 1005)]
ROW_EDGES = [y / 842 for y in (158, 290, 422, 553, 685)]

engine = RapidOCR(params={"Global.log_level": "error"})


def zone(xf, yf):
    col = 1 + sum(xf > e for e in COL_EDGES)
    row = "ABCDEF"[sum(yf > e for e in ROW_EDGES)]
    return f"{row}{col}"


def ocr_tiles(img):
    """Devuelve [(x0,y0,x1,y1,texto,score)] en px de img."""
    H, W = img.shape[:2]
    out = []
    for y in range(0, H, TILE - OVER):
        for x in range(0, W, TILE - OVER):
            t = img[y:y + TILE, x:x + TILE]
            if t.min() > 200:  # tile en blanco
                continue
            r = engine(t)
            if r.boxes is None:
                continue
            for b, txt, s in zip(r.boxes, r.txts, r.scores):
                if s < MIN_SCORE or not txt.strip():
                    continue
                xs, ys = b[:, 0] + x, b[:, 1] + y
                out.append((xs.min(), ys.min(), xs.max(), ys.max(), txt, float(s)))
    return out


def iou(a, b):
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    small = min((a[2] - a[0]) * (a[3] - a[1]), (b[2] - b[0]) * (b[3] - b[1]))
    return inter / small if small else 0


def clean(t):
    return t.replace("十", "+").replace("：", ":").replace("（", "(").replace("）", ")").strip()


def process_page(page):
    pil = page.render(scale=SCALE, grayscale=True).to_pil()
    img = np.array(pil)
    H, W = img.shape
    items = []
    # pasada horizontal
    for x0, y0, x1, y1, t, s in ocr_tiles(img):
        if (x1 - x0) >= (y1 - y0) * 0.6 or len(t) == 1:
            items.append([x0, y0, x1, y1, clean(t), s, 0])
    # pasada vertical (texto que se lee de abajo hacia arriba)
    rot = cv2.rotate(img, cv2.ROTATE_90_CLOCKWISE)
    for x0, y0, x1, y1, t, s in ocr_tiles(rot):
        if len(t) < 2 or (x1 - x0) < (y1 - y0) * 0.6:
            continue
        # rotado (xr,yr) -> original (yr, H-1-xr)
        items.append([y0, H - 1 - x1, y1, H - 1 - x0, clean(t), s, 90])
    # quitar duplicados (solape de tiles / pasadas)
    items.sort(key=lambda i: (-len(i[4]), -i[5]))
    kept = []
    for it in items:
        if all(iou(it, k) < 0.3 for k in kept):
            kept.append(it)
    return kept, W, H


def main(src, outdir):
    base = os.path.splitext(os.path.basename(src))[0]
    doc = pdfium.PdfDocument(src)
    reader = PdfReader(src)
    writer = PdfWriter()
    rows = []
    for pi in range(len(doc)):
        page = doc[pi]
        pw, ph = page.get_size()
        items, W, H = process_page(page)
        print(f"Pagina {pi + 1}/{len(doc)}: {len(items)} textos", flush=True)
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=(pw, ph))
        for x0, y0, x1, y1, t, s, ang in items:
            X0, X1 = x0 / SCALE, x1 / SCALE
            Y0, Y1 = ph - y1 / SCALE, ph - y0 / SCALE  # coords PDF (origen abajo)
            tobj = c.beginText()
            tobj.setTextRenderMode(3)  # invisible
            if ang == 0:
                size = max(Y1 - Y0, 1)
                c.saveState(); c.translate(X0, Y0 + size * 0.15)
                w = c.stringWidth(t, "Helvetica", size) or 1
                c.scale((X1 - X0) / w, 1)
            else:
                size = max(X1 - X0, 1)
                c.saveState(); c.translate(X1 - size * 0.15, Y0); c.rotate(90)
                w = c.stringWidth(t, "Helvetica", size) or 1
                c.scale((Y1 - Y0) / w, 1)
            tobj.setFont("Helvetica", size)
            tobj.setTextOrigin(0, 0)
            tobj.textOut(t)
            c.drawText(tobj)
            c.restoreState()
            cx, cy = (x0 + x1) / 2 / W, (y0 + y1) / 2 / H
            for tok in re.findall(r"[A-Za-z0-9]+", t):
                if any(ch.isdigit() for ch in tok):
                    rows.append((tok, pi + 1, zone(cx, cy), t))
        c.save()
        buf.seek(0)
        overlay = PdfReader(buf).pages[0]
        p = reader.pages[pi]
        p.merge_page(overlay)
        writer.add_page(p)
    for p in writer.pages:
        p.compress_content_streams(level=9)
    writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
    out_pdf = os.path.join(outdir, base + " - BUSCABLE.pdf")
    with open(out_pdf, "wb") as f:
        writer.write(f)
    write_index(rows, os.path.join(outdir, base + " - INDICE.xlsx"))
    print("OK:", out_pdf)


def write_index(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active; ws.title = "Indice"
    ws.append(["Numero / Referencia", "Hojas", "Ubicaciones (hoja:zona)"])
    agg = {}
    for tok, pg, z, _ in rows:
        agg.setdefault(tok, set()).add((pg, z))
    def key(k):
        return (0, int(k), k) if k.isdigit() else (1, 0, k)
    for tok in sorted(agg, key=key):
        locs = sorted(agg[tok])
        ws.append([tok, ", ".join(str(p) for p in sorted({p for p, _ in locs})),
                   ", ".join(f"{p}:{z}" for p, z in locs)])
    ws2 = wb.create_sheet("Detalle")
    ws2.append(["Numero", "Hoja PDF", "Zona", "Texto completo leido"])
    for r in sorted(rows, key=lambda r: (key(r[0]), r[1])):
        ws2.append(list(r))
    for w in (ws, ws2):
        for cell in w[1]:
            cell.font = Font(bold=True)
        w.freeze_panes = "A2"; w.auto_filter.ref = w.dimensions
        w.column_dimensions["A"].width = 22; w.column_dimensions["B"].width = 14
        w.column_dimensions["C"].width = 40; w.column_dimensions["D"].width = 40
    wb.save(path)


if __name__ == "__main__":
    src = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(src))
    main(src, outdir)
