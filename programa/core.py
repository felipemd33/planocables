"""Proceso completo de un plano PDF vectorial:
  - lee el texto de los trazos (sin OCR salvo casos puntuales)
  - reconstruye los cables y asigna numero, color y seccion
  - genera PDF buscable (Ctrl+F) y Excel con el listado de cables
"""
import os, re, io, time, math, collections, datetime
import pypdf
from pdfvec import page_strokes, layer_names
from textdec import Decoder, page_text, UNROT
from wires import WireGraph, assign, NUM_RE, LABEL_RE, norm_color

SKIP_TEXT_LAYERS = ('COMPONENT', 'RIEL', 'TABLERO', 'CABLECANAL', 'Envolvente', 'Zona Segura', 'WATERMARK')
SHREF_RE = re.compile(r'\(\s*Sh\s*(\d+)[A-Z]?\s*[:;.]\s*([A-F]\d)\s*\)', re.I)   # (Sh15A:D2): la letra de subhoja no cuenta, como en instructivo
NOTE_RE = re.compile(r'not\s+indicated\s+will\s+be\s+(\w+)\s+([\d.,]+)\s*mm', re.I)
NOTE_ES_RE = re.compile(r'no\s+indicad\w*\s+(?:ser[aá]n?|son)\s+(\w+)\s+(?:de\s+)?([\d.,]+)\s*mm', re.I)
TAG_RE = re.compile(r'(?!\d+V(?:DC|AC|CC)?$)\d{2}[A-Z][A-Z0-9]{1,7}')   # referencia de componente/bornera: 62XDO, 13XC1, 43DIB1...
TB_RE = re.compile(r'(PAG|REV|GRAL\s*REV|DOC\s*NUMBER|TITLE|CODE|PROJECT|CLIENT|CONT|DATE|DESCRIPTION|PROJ\.?|DW\.?|APR\.?)\s*[:.]?', re.I)
LEGEND_RE = re.compile(r'referencia\s+de\s+s[ií]mbolos|simbolog[ií]a|leyenda|legend', re.I)
COLOR_EN = {'black': 'Negro', 'red': 'Rojo', 'blue': 'Azul', 'white': 'Blanco', 'brown': 'Marrón', 'grey': 'Gris', 'gray': 'Gris'}


# ------------------------------------------------------------------ datos de la hoja
def page_meta(lines, pw, ph, k=1.0, strokes=None):
    """k = escala del dibujo respecto del A1 con que se calibro (texto de los numeros de cable / 7.93 pt)"""
    meta = {}
    # rejilla de zonas: numeros de columna arriba, letras de fila a la izquierda
    cols = sorted(((l['bbox'][0] + l['bbox'][2]) / 2, l['text']) for l in lines
                  if l['text'] in list('123456789') and l['bbox'][1] > 0.88 * ph)
    rows = sorted((-(l['bbox'][1] + l['bbox'][3]) / 2, l['text']) for l in lines
                  if l['text'] in list('ABCDEFGH') and l['bbox'][2] < 0.12 * pw)
    cols = [c for i, c in enumerate(cols) if i == 0 or c[1] != cols[i - 1][1]]
    rows = [r for i, r in enumerate(rows) if i == 0 or r[1] != rows[i - 1][1]]
    if len(cols) >= 3 and [c[1] for c in cols] == sorted(c[1] for c in cols):
        meta['col_edges'] = [((a[0] + b[0]) / 2, b[1]) for a, b in zip(cols, cols[1:])]
        meta['col_first'] = cols[0][1]
    if len(rows) >= 3:
        meta['row_edges'] = [(-(a[0] + b[0]) / 2, b[1]) for a, b in zip(rows, rows[1:])]
        meta['row_first'] = rows[0][1]
    # numero de hoja (PAG) y titulo
    pag = [l for l in lines if re.fullmatch(r'PAG[:.]?|PAG\.?:', l['text'].strip(), re.I)]
    cont = [l for l in lines if re.fullmatch(r'CONT[:.]?', l['text'].strip(), re.I)]
    if pag:
        pag.sort(key=lambda l: (l['bbox'][1] - l['bbox'][0]))   # la del cajetin: abajo a la derecha
        p = pag[0]['bbox']
        cand = [l for l in lines if re.fullmatch(r'\d{1,4}', l['text'].replace(' ', '')) and l['bbox'][3] < p[1] + 1 * k
                and l['bbox'][3] > p[1] - 45 * k and l['bbox'][2] > p[0] - 10 * k and l['bbox'][0] < p[2] + 15 * k]
        if cand:
            # el numero grande puede venir en trozos ('1' + '4'): unir los del mismo renglon
            top = max(l['bbox'][3] for l in cand)
            row = sorted((l for l in cand if abs(l['bbox'][3] - top) < 4 * k), key=lambda l: l['bbox'][0])
            meta['sheet'] = ''.join(l['text'].replace(' ', '') for l in row)
    elif cont:
        # cajetin sin rotulo 'PAG': el numero de hoja es el numero grande de la celda que esta ENCIMA de 'CONT:'
        # (debajo de 'CONT:' / a su derecha va la hoja siguiente)
        cont.sort(key=lambda l: (l['bbox'][1] - l['bbox'][0]))
        p = cont[0]['bbox']; pag = cont
        cand = [l for l in lines if re.fullmatch(r'\d{1,4}', l['text'].replace(' ', '')) and l['bbox'][1] > p[3] - 1 * k
                and l['bbox'][1] < p[3] + 30 * k and l['bbox'][2] > p[0] - 10 * k and l['bbox'][0] < p[2] + 15 * k]
        if cand:
            bot = min(l['bbox'][1] for l in cand)
            row = sorted((l for l in cand if abs(l['bbox'][1] - bot) < 4 * k), key=lambda l: l['bbox'][0])
            meta['sheet'] = ''.join(l['text'].replace(' ', '') for l in row)
    tit = [l for l in lines if l['text'].strip().upper().startswith('TITLE')]
    if tit and pag:
        t = tit[0]['bbox']; p = pag[0]['bbox']
        cand = [l for l in lines if l['bbox'][0] > t[0] - 5 * k and l['bbox'][2] < p[0] - 5 * k and p[1] - 60 * k < l['bbox'][1] < t[3]
                and l is not tit[0] and any(ch.isalpha() for ch in l['text'])]
        cand.sort(key=lambda l: (-l['bbox'][3], l['bbox'][0]))
        meta['title'] = ' '.join(l['text'] for l in cand)
    # cajetin: zona que rodea sus rotulos (sirve aunque el PDF no tenga capas)
    tbl = [l['bbox'] for l in lines if TB_RE.fullmatch(l['text'].strip())]
    if len(tbl) >= 3:
        top = max(b[3] for b in tbl)
        # el borde de arriba del cajetin es la linea larga del marco que pasa justo encima de los rotulos (75287: a 6 pt,
        # 66817: a 3.3 pt); el margen fijo de 25*k llegaba hasta los cables que corren pegados al marco
        from wires import is_wire_layer
        frame = [p[0][1] for layer, op, p in (strokes or ()) if not is_wire_layer(layer) and len(p) >= 2
                 and abs(p[-1][1] - p[0][1]) < 0.3 and abs(p[-1][0] - p[0][0]) > 0.25 * pw and top - 1 < p[0][1] < top + 25 * k]
        meta['cajetin'] = (min(b[0] for b in tbl) - 30 * k, min(b[1] for b in tbl) - 40 * k,
                           max(b[2] for b in tbl) + 110 * k, (min(frame) + 0.5) if frame else top + 25 * k)
    for l in lines:
        m = NOTE_RE.search(l['text']) or NOTE_ES_RE.search(l['text'])
        if m:
            c = m.group(1).lower()
            meta['default'] = (COLOR_EN.get(c, norm_color(c)), m.group(2).replace('.', ','), l['text'])
    return meta


def zone_of(meta, x, y, pw, ph):
    if 'col_edges' in meta:
        col = meta['col_first']
        for e, lab in meta['col_edges']:
            if x > e: col = lab
    else:
        col = str(1 + min(7, int((x / pw - 0.08) / 0.105))) if x / pw > 0.08 else '1'
    if 'row_edges' in meta:
        row = meta['row_first']
        for e, lab in meta['row_edges']:
            if y < e: row = lab
    else:
        row = 'ABCDEF'[max(0, min(5, int((1 - y / ph - 0.04) / 0.157)))]
    return f'{row}{col}'


# ------------------------------------------------------------------ proceso
class Result:
    pass


def process(pdf_path, log=print, use_ocr=True, pages=None):
    import eplan
    if eplan.es_eplan(pdf_path):      # plano de EPLAN (texto real + lista de conexiones): otro lector, misma interfaz
        return eplan.process(pdf_path, log=log, use_ocr=use_ocr, pages=pages)
    t0 = time.time()
    reader = pypdf.PdfReader(pdf_path)
    names = layer_names(reader)
    dec = Decoder(use_ocr=use_ocr)
    res = Result(); res.pages = []; res.path = pdf_path
    n = len(reader.pages)
    leidas = []
    for pi in range(n):
        if pages and pi + 1 not in pages:
            continue
        st = page_strokes(reader, pi, names)
        todas = [l for l in page_text(st, dec, SKIP_TEXT_LAYERS) if l['text'].strip()]
        lines = [l for l in todas if not set(l['text']) <= set('?_-•. ')]
        # el '-' suelto no entra en los textos (casi siempre es una raya de corte o de una discontinua), pero se guarda
        # aparte: puede ser el rotulo del polo '-' de un aparato (pin de un cargador, borne de una bateria)
        signos = [l for l in todas if l['text'].strip() == '-']
        leidas.append((pi, st, lines, signos))
        log(f'Leyendo textos de la hoja {pi + 1}/{n}…')
    # escala del dibujo: altura del texto de los numeros de cable (7.93 pt en el A1 con que se calibro; ~4.3 en un A3)
    H = text_scale([l for _, _, ls, _ in leidas for l in ls])
    k = H / H_REF
    res.H, res.k = H, k
    for pi, st, lines, signos in leidas:
        mb = reader.pages[pi].mediabox
        pw, ph = float(mb.right), float(mb.top)
        mx0, my0 = float(mb.left), float(mb.bottom)
        meta = page_meta(lines, pw, ph, k, st)
        g = WireGraph(st, lines, H=H)
        tb = meta.get('cajetin')
        in_tb = lambda l: tb and tb[0] <= (l['bbox'][0] + l['bbox'][2]) / 2 <= tb[2] and tb[1] <= (l['bbox'][1] + l['bbox'][3]) / 2 <= tb[3]
        nums, labels = assign(g, [l for l in lines if not in_tb(l)], H=H)   # nada del cajetin es un cable
        keep = g.dashed_wanted(nums, labels)
        if keep:   # cables dibujados en linea discontinua (cableado opcional): se rearma la hoja conservandolos
            g = WireGraph(st, lines, H=H, keep_dashed=keep)
            lines = [l for l in lines if not g.dash_text(l)]   # 'letras' que eran rayas de la discontinua
            nums, labels = assign(g, [l for l in lines if not in_tb(l)], H=H)
        if LEGEND_RE.search(meta.get('title', '')):
            nums, labels = [], []   # hoja de leyenda: los numeros son ejemplos
        pg = dict(index=pi + 1, w=pw, h=ph, x0=mx0, y0=my0, lines=lines, signos=signos, strokes=st, meta=meta, graph=g, nums=nums, labels=labels,
                  raster=(not st and has_images(reader.pages[pi])))
        res.pages.append(pg)
        log(f'Hoja {pi + 1}/{n}: {len(lines)} textos, {sum(1 for x in nums if x["chain"] is not None)} números de cable')
    dec.save_cache()
    res.default = next((p['meta']['default'] for p in res.pages if 'default' in p['meta']), None)
    res.cables, res.detail, res.review = build_cable_list(res)
    res.stats = dict(dec.stats); res.seconds = time.time() - t0
    return res


H_REF = 7.93   # altura del texto de los numeros de cable en el plano con que se calibraron los umbrales (75287, A1)


def text_scale(lines):
    """altura tipica del texto de los numeros de cable del plano (mediana de los renglones que son solo un numero
    de 3-5 cifras en capas de cable). Dentro de +-10% del plano de referencia se usa la referencia tal cual."""
    from wires import NON_WIRE_LAYERS
    hs = sorted(l['H'] for l in lines if re.fullmatch(r'\d{3,5}', l['text'].strip()) and l.get('layer') not in NON_WIRE_LAYERS)
    if len(hs) < 5:
        return H_REF
    h = hs[len(hs) // 2]
    return H_REF if abs(h / H_REF - 1) < 0.10 else h


def has_images(page):
    try:
        return bool((page.get('/Resources') or {}).get('/XObject'))
    except Exception:
        return False


def sheet_name(pg):
    return pg['meta'].get('sheet', str(pg['index']))


def build_cable_list(res):
    detail = []; review = []
    chain_labels = {}
    for pg in res.pages:
        for lb in pg['labels']:
            if lb['chain'] is not None:
                chain_labels.setdefault((pg['index'], lb['chain']), []).append(lb)
            else:
                bb = lb['bbox']
                review.append(dict(tipo='Etiqueta de color/sección sin cable asociado', texto=lb['raw'], pag=pg['index'], bbox=bb,
                                   hoja=sheet_name(pg), zona=zone_of(pg['meta'], (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2, pg['w'], pg['h'])))
    for pg in res.pages:
        g = pg['graph']
        for nm in pg['nums']:
            bb = nm['bbox']; z = zone_of(pg['meta'], (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2, pg['w'], pg['h'])
            if nm['chain'] is None:
                if nm['whole'] and len(nm['num']) == 4 and nm['line'].get('layer') not in ('Texto Rotulo',):
                    review.append(dict(tipo='Número sin cable asociado (¿es un cable?)', texto=nm['line']['text'],
                                       pag=pg['index'], hoja=sheet_name(pg), zona=z, bbox=bb))
                continue
            lbs = chain_labels.get((pg['index'], nm['chain']), [])
            via = ''
            if not lbs:   # tramo sin etiqueta: buscar en el cable que sigue recto tras una union
                seen = {nm['chain']}; frontier = [nm['chain']]
                for _ in range(4):
                    nxt = []
                    for c in frontier:
                        for c2 in g.straight_neighbors(c):
                            if c2 not in seen:
                                seen.add(c2); nxt.append(c2)
                    found = [lb for c2 in nxt for lb in chain_labels.get((pg['index'], c2), [])]
                    if found:
                        lbs = found; via = 'recto'; break
                    frontier = nxt
            refs = sheet_refs(pg, g.chain_ends(nm['chain']), res.k)
            detail.append(dict(num=nm['num'], pag=pg['index'], hoja=sheet_name(pg), zona=z, bbox=bb,
                               etiquetas=[(lb['color'], lb['sec']) for lb in lbs], via=via, corte='tick' in nm,
                               refs=refs, texto=nm['line']['text'], chain=(pg['index'], nm['chain']),
                               src=nm['line'].get('src', '')))
    # color/seccion por aparicion: tramo propio -> otra aparicion del mismo numero -> nota del plano
    by_num = collections.defaultdict(list)
    for d in detail:
        by_num[d['num']].append(d)
    for num, occ in by_num.items():
        own = collections.Counter(e for d in occ if not d['via'] for e in d['etiquetas'])
        if not own:
            own = collections.Counter(e for d in occ for e in d['etiquetas'])
        for d in occ:
            if d['etiquetas']:
                d['color'], d['sec'] = d['etiquetas'][0]
                d['origen'] = 'Etiqueta en el mismo tramo' if not d['via'] else 'Etiqueta del mismo cable (continúa recto tras una unión)'
                if len(set(d['etiquetas'])) > 1:
                    d['origen'] += ' (hay varias: ' + ', '.join(f'{c} {s}' for c, s in sorted(set(d['etiquetas']))) + ')'
            elif own:
                (c, s), _ = own.most_common(1)[0]
                src_pg = sorted({x['hoja'] for x in occ if (c, s) in x['etiquetas']}, key=natkey)
                d['color'], d['sec'] = c, s
                d['origen'] = 'Etiqueta en otra aparición (hoja ' + ', '.join(src_pg) + ')'
            elif res.default:
                d['color'], d['sec'] = res.default[0], res.default[1]
                d['origen'] = 'Sin indicar en el plano → nota: ' + res.default[0] + ' ' + res.default[1] + ' mm²'
            else:
                d['color'], d['sec'] = '', ''
                d['origen'] = 'Sin indicar en el plano'
    build_routes(res, detail, chain_labels)
    route_refs(res, detail)
    for d in detail:
        d['puntas'] = res.pages_by_index[d['pag']]['routes'].get(d['num'], {}).get('puntas', 0)
    cables = []
    for num in sorted(by_num, key=natkey):
        occ = by_num[num]
        combos = collections.OrderedDict()
        for d in sorted(occ, key=lambda d: (natkey(d['hoja']), d['zona'])):
            combos.setdefault((d['color'], d['sec']), []).append(d)
        for (c, s), ds in combos.items():
            obs = []
            if len(combos) > 1:
                obs.append('Este número tiene tramos con distinto color/sección: revisar')
            if all(x['origen'].startswith('Sin indicar') for x in ds):
                obs.append('Color y sección según la nota del plano')
            elif any(x['origen'].startswith('Etiqueta en otra') for x in ds) and not any(x['origen'].startswith('Etiqueta en el mismo') for x in ds):
                obs.append(ds[0]['origen'])
            if any(x['origen'].startswith('Etiqueta del mismo cable') for x in ds):
                obs.append('Color/sección tomados del tramo recto contiguo')
            if any('hay varias' in x['origen'] for x in ds):
                obs.append('Un mismo tramo tiene dos etiquetas distintas: revisar')
            der = sorted({(x['hoja'], x['pag'], x['puntas']) for x in ds if x['puntas'] >= 3 and res.pages_by_index[x['pag']]['routes'][num]['uniones']},
                         key=lambda t: natkey(t[0]))
            for hoja, pag, n in der:
                junto = res.pages_by_index[pag]['routes'][num].get('junto') or []
                obs.insert(0, f'Cable de {n} puntas en hoja {hoja}: puente/derivación' + (f' junto a {", ".join(junto)}' if junto else '')
                           + ' (hecho a propósito)')
            if any(res.pages_by_index[x['pag']]['routes'].get(num, {}).get('discontinuo') for x in ds):
                obs.append('Tramo dibujado en línea discontinua (cableado opcional): revisar si va')
            puntas = max((x['puntas'] for x in ds), default=0)
            cables.append(dict(num=num, color=c, sec=s, puntas=puntas,
                               hojas=', '.join(sorted({x['hoja'] for x in ds}, key=natkey)),
                               ubic=', '.join(f"{x['hoja']}:{x['zona']}" for x in ds),
                               n=len(ds), refs=', '.join(sorted({r for x in ds for r in x['refs']})),
                               obs='; '.join(obs)))
            if len(combos) > 1 or any('hay varias' in x['origen'] for x in ds):
                pass
    numbered = {d['chain'] for d in detail}
    res.unnumbered = []
    for pg in res.pages:
        for lb in pg['labels']:
            if lb['chain'] is None or (pg['index'], lb['chain']) in numbered:
                continue
            g = pg['graph']
            net_chains = [c for c in range(len(g.chains)) if g.chain_net[c] == g.chain_net[lb['chain']]]
            if any((pg['index'], c) in numbered and c in g.straight_neighbors(lb['chain']) for c in net_chains):
                continue   # es un tramo de un cable numerado que sigue recto
            if any(lb['chain'] in r.get('chains', ()) for r in pg.get('routes', {}).values()):
                continue   # es parte del recorrido de un cable numerado (p.ej. tras el salto de una discontinua)
            bb = lb['bbox']
            refs = sheet_refs(pg, g.chain_ends(lb['chain']), res.k)
            res.unnumbered.append(dict(color=lb['color'], sec=lb['sec'], hoja=sheet_name(pg), pag=pg['index'], bbox=bb,
                                       zona=zone_of(pg['meta'], (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2, pg['w'], pg['h']),
                                       largo=round(g.chain_length(lb['chain'])), refs=', '.join(refs), texto=lb['raw']))
    for num, occ in by_num.items():
        combos = {(d['color'], d['sec']) for d in occ}
        if len(combos) > 1:
            review.append(dict(tipo='Cable con distintos colores/secciones', texto=num + ': ' + ' | '.join(f'{c} {s}' for c, s in sorted(combos)),
                               pag=occ[0]['pag'], hoja=occ[0]['hoja'], zona=occ[0]['zona'], bbox=occ[0]['bbox'], num=num))
    return cables, detail, review


def build_routes(res, detail, chain_labels):
    """recorrido de cada numero de cable en cada hoja, con el color de cada tramo"""
    res.pages_by_index = {p['index']: p for p in res.pages}
    by_page = collections.defaultdict(lambda: collections.defaultdict(list))
    for d in detail:
        by_page[d['pag']][d['num']].append(d)
    for pg in res.pages:
        pg['routes'] = {}
        g = pg['graph']
        occ = by_page.get(pg['index'], {})
        tags = [l for l in pg['lines'] if TAG_RE.fullmatch(l['text'].strip())]
        numbered = collections.defaultdict(set)   # tramo -> numeros
        for num, ds in occ.items():
            for d in ds:
                numbered[d['chain'][1]].add(num)
        labeled = {c: {(lb['color'], lb['sec']) for lb in lbs} for (pi, c), lbs in chain_labels.items() if pi == pg['index']}
        for num, ds in occ.items():
            seeds = {d['chain'][1] for d in ds}
            mine = {(d['color'], d['sec']) for d in ds}
            blocked = {c for c, ns in numbered.items() if num not in ns}
            # un tramo sin numero con etiqueta propia distinta es otro conductor (p.ej. tierra A-V/6mm2)
            blocked |= {c for c, labs in labeled.items() if c not in numbered and not (labs & mine)}
            chains, ends, joins = g.route(seeds, blocked)
            base = collections.Counter((d['color'], d['sec']) for d in ds).most_common(1)[0][0]
            # color de cada tramo: su etiqueta, o la del tramo numerado mas cercano del recorrido
            own = {c: (lb[0]['color'], lb[0]['sec']) for c in chains for lb in [chain_labels.get((pg['index'], c), [])] if lb}
            for d in ds:
                own.setdefault(d['chain'][1], (d['color'], d['sec']))
            col = dict(own)
            frontier = list(own)
            while frontier:
                nxt = []
                for c in frontier:
                    for k in g.chain_end_keys(c):
                        for _, s2 in g.adj.get(k, []):
                            c2 = g.seg_chain.get(s2)
                            if c2 in chains and c2 not in col:
                                col[c2] = col[c]; nxt.append(c2)
                frontier = nxt
            tramos = []
            for c in sorted(chains):
                color, sec = col.get(c, base)
                segs = [[round(v, 1) for v in (*g.segs[sid][0], *g.segs[sid][1])] for sid in g.chains[c]]
                tramos.append(dict(color=color, sec=sec, segs=segs))
            junto = []
            for j in joins:   # bornera/componente mas cercano a cada union
                best = min(((math.hypot((l['bbox'][0] + l['bbox'][2]) / 2 - j[0], (l['bbox'][1] + l['bbox'][3]) / 2 - j[1]), l['text'])
                            for l in tags), default=None)
                if best and best[0] < 160 * res.k and best[1] not in junto:
                    junto.append(best[1])
            pg['routes'][num] = dict(tramos=tramos, puntas=len(ends), junto=junto, chains=sorted(chains),
                                     discontinuo=bool(chains & getattr(g, 'dashed_chains', set())),
                                     fines=[[round(p[0], 1), round(p[1], 1)] for p in ends],
                                     uniones=[[round(p[0], 1), round(p[1], 1)] for p in joins])


def sheet_refs(pg, ends, k):
    """referencias '(Sh62:D6)' de las puntas de un tramo: rotulos a menos de 45 pt (escalados por k) de una punta"""
    out = set()
    for e in ends:
        for l in pg['lines']:
            if not near(l['bbox'], e, 45 * k):
                continue
            for x in [l]:   # (probado: sumar el rotulo apilado entero trae referencias de flechas vecinas)
                for m in SHREF_RE.finditer(x['text']):
                    out.add(f'Sh{m.group(1)}:{m.group(2).upper()}')
    return sorted(out)


def stacked_ref(lines, p, k):
    """(ShNN:XX) al pie del rotulo de varios renglones centrado en la punta p (flecha): 'Comand' / 'Timer Relay' /
    '-0Vdc' / '(Sh62:D3)'. Se baja (o sube) renglon por renglon desde el mas cercano (como instructivo.ref_apilada)"""
    horiz = [l for l in lines if l.get('ang', 0) in (0, 180)]
    centrado = lambda l: l['bbox'][0] - 2 * k <= p[0] <= l['bbox'][2] + 2 * k
    first = [l for l in horiz if centrado(l) and near(l['bbox'], p, 30 * k)]
    if not first:
        return None
    cur = min(first, key=lambda l: math.hypot(max(l['bbox'][0] - p[0], 0, p[0] - l['bbox'][2]), max(l['bbox'][1] - p[1], 0, p[1] - l['bbox'][3])))
    abajo = (cur['bbox'][1] + cur['bbox'][3]) / 2 < p[1]
    seen = [cur]
    for _ in range(8):
        if SHREF_RE.search(cur['text']):
            return cur
        h = cur['bbox'][3] - cur['bbox'][1]
        gap = (lambda l: cur['bbox'][1] - l['bbox'][3]) if abajo else (lambda l: l['bbox'][1] - cur['bbox'][3])
        nxt = [l for l in horiz if all(l is not x for x in seen) and centrado(l) and -0.3 * h <= gap(l) < 1.2 * h]
        if not nxt:
            break
        cur = min(nxt, key=gap); seen.append(cur)
    return None


def route_refs(res, detail):
    """'Continua en' por recorrido: en cada punta del recorrido del numero en la hoja, la referencia del rotulo apilado
    de su flecha; si no hay pila, los rotulos a menos de 45*k (como antes) que no son la pila de otra punta de la hoja
    (66817 hoja 12: 1215 se llevaba el Sh62:D3 de la flecha de 1216, y 1216 perdia la suya por estar 4 renglones abajo)"""
    k = res.k
    for pg in res.pages:
        rts = pg.get('routes') or {}
        pila = {}
        for num, rt in rts.items():
            for e in rt.get('fines') or []:
                l = stacked_ref(pg['lines'], e, k)
                if l is not None:
                    pila[(num, tuple(e))] = l
        claimed = {id(l) for l in pila.values()}
        por_num = {}
        for num, rt in rts.items():
            out = set()
            for e in rt.get('fines') or []:
                l = pila.get((num, tuple(e)))
                cand = [l] if l is not None else [x for x in pg['lines'] if id(x) not in claimed and near(x['bbox'], e, 45 * k)]
                for x in cand:
                    out.update(f'Sh{m.group(1)}:{m.group(2).upper()}' for m in SHREF_RE.finditer(x['text']))
            por_num[num] = sorted(out)
        for d in detail:
            if d['pag'] == pg['index'] and d['num'] in por_num:
                d['refs'] = por_num[d['num']]


def near(bb, p, r):
    dx = max(bb[0] - p[0], 0, p[0] - bb[2]); dy = max(bb[1] - p[1], 0, p[1] - bb[3])
    return math.hypot(dx, dy) <= r


def natkey(s):
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', str(s))]


# ------------------------------------------------------------------ PDF buscable
def write_searchable_pdf(res, out_path, raster_ocr=None):
    from reportlab.pdfgen import canvas
    reader = pypdf.PdfReader(res.path)
    writer = pypdf.PdfWriter()
    done = {pg['index']: pg for pg in res.pages}
    for pi, page in enumerate(reader.pages):
        pg = done.get(pi + 1)
        items = []
        if pg:
            items = [(l['text'], l['bbox'], l['ang'], l.get('base'), l.get('x0'), l.get('x1'), l.get('H')) for l in pg['lines']]
        if pg and pg.get('raster') and raster_ocr:
            items += raster_ocr(res.path, pi)
        if items:
            pw, ph = [float(v) for v in page.mediabox[2:]]
            buf = io.BytesIO(); c = canvas.Canvas(buf, pagesize=(pw, ph))
            for it in items:
                put_text(c, *it)
            c.save(); buf.seek(0)
            page.merge_page(pypdf.PdfReader(buf).pages[0])
        writer.add_page(page)
    for p in writer.pages:
        p.compress_content_streams(level=9)
    writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
    with open(out_path, 'wb') as f:
        writer.write(f)


def safe_text(t):
    out = []
    for ch in t:
        try:
            ch.encode('cp1252'); out.append(ch)
        except UnicodeEncodeError:
            out.append('?')
    return ''.join(out)


def put_text(c, text, bb, ang, base=None, x0=None, x1=None, H=None):
    text = safe_text(text)
    if not text.strip():
        return
    if base is None or H is None:  # solo caja (OCR de imagen)
        x0, y0, x1b, y1 = bb; ang = ang or 0; H = (y1 - y0) if ang in (0, 180) else (x1b - x0)
        base = y0; x1 = x1b if ang in (0, 180) else y1
        ox, oy = (x0, y0) if ang == 0 else (x1b, y0)
    else:
        ox, oy = UNROT[ang](x0, base)
    size = max(H / 0.72, 1.0)
    width = max((x1 - x0) if x1 is not None else 1, 0.5)
    w = c.stringWidth(text, 'Helvetica', size) or 1
    c.saveState()
    c.translate(ox, oy)
    c.rotate({0: 0, 90: 90, 270: -90, 180: 180}.get(ang, 0))
    c.scale(width / w, 1)
    t = c.beginText(); t.setTextRenderMode(3); t.setFont('Helvetica', size); t.setTextOrigin(0, 0)
    t.textOut(text); c.drawText(t)
    c.restoreState()


# ------------------------------------------------------------------ Excel
def write_excel(res, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F4E78')
    warn = PatternFill('solid', fgColor='FFF2CC'); note = PatternFill('solid', fgColor='E2EFDA')

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font = head; c.fill = fill; c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30

    ws = wb.active; ws.title = 'Listado de cables'
    sheet(ws, ['Nº cable', 'Color', 'Sección (mm²)', 'Hojas del plano', 'Ubicación (hoja:zona)', 'Nº etiquetas',
               'Puntas', 'Continúa en', 'Observaciones'],
          [[natnum(c['num']), c['color'], secnum(c['sec']), c['hojas'], c['ubic'], c['n'], c.get('puntas') or None, c['refs'], c['obs']] for c in res.cables],
          [11, 16, 13, 16, 40, 11, 9, 18, 60])
    blue = PatternFill('solid', fgColor='DDEBF7')
    for row in ws.iter_rows(min_row=2):
        obs = row[8].value or ''
        if 'puntas' in obs:
            row[6].fill = blue; row[8].fill = blue
        if 'revisar' in obs:
            for c in row: c.fill = warn
        elif 'nota del plano' in obs:
            for c in row[1:3]: c.fill = note

    ws2 = wb.create_sheet('Detalle por aparición')
    sheet(ws2, ['Nº cable', 'Hoja del plano', 'Página PDF', 'Zona', 'Color', 'Sección (mm²)', 'Origen del color/sección',
                'Raya de corte', 'Continúa en', 'Texto leído'],
          [[natnum(d['num']), natnum(d['hoja']), d['pag'], d['zona'], d['color'], secnum(d['sec']), d['origen'],
            'Sí' if d['corte'] else 'No', ', '.join(d['refs']), d['texto']]
           for d in sorted(res.detail, key=lambda d: (natkey(d['num']), natkey(d['hoja']), d['zona']))],
          [11, 13, 11, 8, 16, 13, 50, 12, 18, 30])

    wsu = wb.create_sheet('Cables sin número')
    sheet(wsu, ['Color', 'Sección (mm²)', 'Hoja del plano', 'Página PDF', 'Zona', 'Continúa en', 'Etiqueta leída'],
          [[u['color'], secnum(u['sec']), natnum(u['hoja']), u['pag'], u['zona'], u['refs'], u['texto']]
           for u in sorted(getattr(res, 'unnumbered', []), key=lambda u: (natkey(u['hoja']), u['zona']))],
          [16, 13, 13, 11, 8, 18, 30])

    ws3 = wb.create_sheet('A revisar')
    sheet(ws3, ['Tipo', 'Texto', 'Hoja del plano', 'Página PDF', 'Zona'],
          [[r['tipo'], r['texto'], natnum(r['hoja']), r['pag'], r['zona']] for r in res.review], [48, 40, 13, 11, 8])

    ws4 = wb.create_sheet('Textos del plano')
    rows = []
    for pg in res.pages:
        for l in sorted(pg['lines'], key=lambda l: (-l['bbox'][3], l['bbox'][0])):
            bb = l['bbox']
            rows.append([natnum(sheet_name(pg)), pg['index'], zone_of(pg['meta'], (bb[0] + bb[2]) / 2, (bb[1] + bb[3]) / 2, pg['w'], pg['h']), l['text']])
    sheet(ws4, ['Hoja del plano', 'Página PDF', 'Zona', 'Texto'], rows, [13, 11, 8, 80])

    ws5 = wb.create_sheet('Resumen')
    total_occ = len(res.detail)
    info = [['Archivo', os.path.basename(res.path)],
            ['Fecha de proceso', datetime.datetime.now().strftime('%d/%m/%Y %H:%M')],
            ['Páginas procesadas', len(res.pages)],
            ['Números de cable distintos', len({c['num'] for c in res.cables})],
            ['Etiquetas de número encontradas', total_occ],
            ['Tramos con etiqueta de color pero sin número', len(getattr(res, 'unnumbered', []))],
            ['Regla por defecto (nota del plano)', res.default[2] if res.default else 'No encontrada'],
            ['Elementos a revisar', len(res.review)],
            ['Tiempo de proceso (s)', round(res.seconds, 1)],
            ['Hojas', ', '.join(f"{sheet_name(p)} ({p['meta'].get('title', '')})" for p in res.pages if p['meta'].get('title'))]]
    for r in info:
        ws5.append(r)
    ws5.column_dimensions['A'].width = 34; ws5.column_dimensions['B'].width = 110
    for r in ws5.iter_rows():
        r[0].font = Font(bold=True)
    wb.move_sheet('Resumen', offset=-5)
    wb.active = 1
    wb.save(out_path)


def natnum(s):
    # '03' se deja como texto para que coincida con el cajetin
    return int(s) if isinstance(s, str) and s.isdigit() and not (len(s) > 1 and s.startswith('0')) else s


def secnum(s):
    try:
        return float(s.replace(',', '.'))
    except Exception:
        return s
